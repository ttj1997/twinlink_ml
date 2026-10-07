"""Three-way comparison document / workbook / rebuild, and the findings about ``Model_v4.1.xlsx``.

The workbook is confidential partner material and is read from a path; this module contains
no workbook content apart from cell addresses and the structure of formulas.

* :func:`compare_sheet` -- per numeric formula cell: Excel's stored value, the as-built
  reproduction (:mod:`.as_built`) and the value following the equation document
  (:mod:`.cascade`); with the cause where they differ,
* :func:`gas_hand_calculation` -- the gas formula decomposed so that it can be redone by hand,
* :func:`inputs_diff`, :func:`reference_scan` -- facts about the workbook (changed inputs,
  unused cells, the two balance sheets),
* :func:`negative_recovery_scan`, :func:`equal_outlet_vs_counter_current` -- the analysis of
  the equal-outlet-temperature assumption (:mod:`.loops`).
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path

from twinlink_ml.partner_v4 import corrected as c

from . import cascade as k
from . import loops
from .inputs import ExcelInputsV41, SiInputsV41

__all__ = [
    "DOCUMENT_TYPOS",
    "EQUATION_TO_CELLS",
    "CellComparison",
    "break_even_htm_flow",
    "compare_sheet",
    "equal_outlet_vs_counter_current",
    "gas_hand_calculation",
    "inputs_diff",
    "negative_recovery_scan",
    "read_formula_cells",
    "reference_scan",
]

_GAS_ROWS = {26, 39}
_TOTAL_ROW = 41


def read_formula_cells(path: str | Path, sheet: str) -> dict[str, tuple[str, float]]:
    """Numeric formula cells of a sheet: ``{address: (formula text, value Excel stored)}``."""
    import openpyxl

    formulas = openpyxl.load_workbook(Path(path))[sheet]
    cached = openpyxl.load_workbook(Path(path), data_only=True)[sheet]
    out: dict[str, tuple[str, float]] = {}
    for row in formulas.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                value = cached[cell.coordinate].value
                if isinstance(value, int | float):
                    out[cell.coordinate] = (cell.value, float(value))
    return out


@dataclass(frozen=True)
class CellComparison:
    """One cell: Excel, as-built reproduction, document-based value."""

    sheet: str
    cell: str
    formula: str
    excel: float
    as_built: float | None
    document: float | None
    cause: str

    @property
    def as_built_vs_excel(self) -> float | None:
        """Relative deviation of the reproduction from Excel (reproduction error)."""
        return _rel(self.as_built, self.excel)

    @property
    def document_vs_excel(self) -> float | None:
        """Excel divided by the document value (the factor by which the workbook is off)."""
        if self.document in (None, 0.0):
            return None
        return self.excel / self.document

    def as_dict(self) -> dict[str, object]:
        """Plain dict for the JSON artifact."""
        d = asdict(self)
        d["as_built_vs_excel"] = self.as_built_vs_excel
        d["document_vs_excel"] = self.document_vs_excel
        return d


def _rel(value: float | None, reference: float) -> float | None:
    if value is None:
        return None
    return abs(value - reference) / max(abs(reference), 1e-300)


def _cause(sheet: str, cell: str) -> str:
    row = int(re.sub(r"[A-Z]+", "", cell))
    if sheet == "Scenarios" and row in _GAS_ROWS:
        return "G: gas formula (efficiency /1000 and moisture in g/kg), factor 108.9"
    if sheet == "Scenarios" and row == _TOTAL_ROW:
        return "G: sum of two gas cells"
    return ""


def compare_sheet(
    sheet: str,
    excel: dict[str, tuple[str, float]],
    as_built: dict[str, float],
    document: dict[str, float],
) -> list[CellComparison]:
    """Compare every numeric formula cell of ``sheet`` (union of the three sources)."""

    def key(address: str) -> tuple[str, int]:
        return re.sub(r"\d+", "", address), int(re.sub(r"[A-Z]+", "", address))

    out = []
    for address in sorted(set(excel) | set(as_built) | set(document), key=key):
        formula, value = excel.get(address, ("", float("nan")))
        out.append(
            CellComparison(
                sheet=sheet,
                cell=address,
                formula=formula,
                excel=value,
                as_built=as_built.get(address),
                document=document.get(address),
                cause=_cause(sheet, address),
            )
        )
    return out


def gas_hand_calculation(
    inputs: ExcelInputsV41, air_kg_per_h: float, t_in: float, t_out: float
) -> dict[str, float | str]:
    """The gas formula of ``Scenarios`` rows 26/39 against (3.2), decomposed for a hand check."""
    from .as_built import gas_kg_per_h_as_built

    m = inputs.plant.material
    y_g = inputs.plant.dryer_1.air_moisture_g_per_kg
    built = gas_kg_per_h_as_built(air_kg_per_h, t_in, t_out, y_g, m)
    si = inputs.to_si()
    stream = c.Stream(air_kg_per_h / 3600.0, y_g / 1e3, t_in)
    right = loops.heater_gas_kg_per_s(stream, t_out, si.plant.material) * 3600.0
    cp_built = m.cp_gas_kj_per_kg_k + y_g * m.cp_vapour_kj_per_kg_k
    cp_right = m.cp_gas_kj_per_kg_k + y_g / 1e3 * m.cp_vapour_kj_per_kg_k
    eta_factor = (m.heater_efficiency_percent / 100.0) / (m.heater_efficiency_percent / 1000.0)
    dt = t_out - t_in
    return {
        "as_built_kg_per_h": built,
        "document_kg_per_h": right,
        "ratio": built / right,
        "moisture_factor": cp_built / cp_right,
        "efficiency_factor": eta_factor,
        "hand_as_built": (
            f"{air_kg_per_h:g}*{dt:g}*{cp_built:.4g} / ({m.calorific_value_kj_per_kg:g}*"
            f"{m.heater_efficiency_percent:g}/1000) = {air_kg_per_h * dt * cp_built:.6g} / "
            f"{m.calorific_value_kj_per_kg * m.heater_efficiency_percent / 1000:.6g}"
        ),
        "hand_document": (
            f"{air_kg_per_h:g}*{dt:g}*{cp_right:.4g} / ({m.calorific_value_kj_per_kg:g}*"
            f"{m.heater_efficiency_percent / 100:g}) = {air_kg_per_h * dt * cp_right:.6g} / "
            f"{m.calorific_value_kj_per_kg * m.heater_efficiency_percent / 100:.6g}"
        ),
    }


def inputs_diff(old: str | Path, new: str | Path) -> list[dict[str, object]]:
    """Cells of the sheet ``Inputs`` that differ between two workbook versions."""
    import openpyxl

    a = openpyxl.load_workbook(Path(old), data_only=True)["Inputs"]
    b = openpyxl.load_workbook(Path(new), data_only=True)["Inputs"]
    addresses = {c_.coordinate for ws in (a, b) for row in ws.iter_rows() for c_ in row}
    out = []
    for address in sorted(addresses, key=lambda x: (int(re.sub(r"[A-Z]+", "", x)), x)):
        va, vb = a[address].value, b[address].value
        if va != vb:
            label = (
                b[f"C{re.sub(r'[A-Z]+', '', address)}"].value
                or b[f"G{re.sub(r'[A-Z]+', '', address)}"].value
            )
            out.append({"cell": address, "label": label, "old": va, "new": vb})
    return out


def reference_scan(path: str | Path) -> dict[str, object]:
    """Facts about references in the workbook (formula text only).

    * cells that refer to ``Inputs`` column ``K`` (expected: none),
    * cells that refer to the sheet ``Balance`` (singular),
    * whether ``Balance`` repeats ``Balances!A1:F15`` in formula text and stored value,
    * the sheets and the number of formulas per sheet.
    """
    import openpyxl

    formulas = openpyxl.load_workbook(Path(path))
    cached = openpyxl.load_workbook(Path(path), data_only=True)
    k_refs, balance_refs = [], []
    per_sheet = {}
    for ws in formulas:
        count = 0
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    count += 1
                    if re.search(r"Inputs!\$?K\$?\d", cell.value):
                        k_refs.append(f"{ws.title}!{cell.coordinate}")
                    if re.search(r"\bBalance!", cell.value):
                        balance_refs.append(f"{ws.title}!{cell.coordinate}")
        per_sheet[ws.title] = count
    same = True
    if "Balance" in formulas.sheetnames and "Balances" in formulas.sheetnames:
        for row in formulas["Balance"].iter_rows(min_row=1, max_row=15, max_col=6):
            for cell in row:
                other = formulas["Balances"][cell.coordinate]
                if cell.value != other.value:
                    same = False
                if (
                    cached["Balance"][cell.coordinate].value
                    != cached["Balances"][cell.coordinate].value
                ):
                    same = False
    return {
        "sheets": formulas.sheetnames,
        "formulas_per_sheet": per_sheet,
        "inputs_column_k_referenced_by": k_refs,
        "balance_sheet_referenced_by": balance_refs,
        "balance_repeats_balances_a1_f15": same,
    }


def _min_stage(result: k.CaseResult) -> tuple[float, int]:
    stage = min(result.stages, key=lambda s: s.recovered_w)
    return stage.recovered_w, stage.source


def break_even_htm_flow(
    inputs: ExcelInputsV41, case: int, order: k.Order | None = None, *, upper_kg_per_h: float = 1e6
) -> float | None:
    """Circuit flow (both circuits equal) above which a stage recovers negative heat.

    Bisection on the smallest stage energy; ``None`` if it stays non-negative up to ``upper``.
    """

    def negative(flow: float) -> bool:
        si = inputs.with_htm_flow((flow, flow)).to_si()
        return _min_energy(si, case, order) < 0.0

    if not negative(upper_kg_per_h):
        return None
    lo, hi = 1e-3, upper_kg_per_h
    for _ in range(200):
        mid = 0.5 * (lo + hi)
        lo, hi = (lo, mid) if negative(mid) else (mid, hi)
    return 0.5 * (lo + hi)


def _min_energy(si: SiInputsV41, case: int, order: k.Order | None) -> float:
    return _min_stage(k.evaluate_case(si, case, order=order))[0]


def negative_recovery_scan(
    inputs: ExcelInputsV41, flows_kg_per_h: tuple[float, ...]
) -> list[dict[str, object]]:
    """Smallest stage energy per case, for each circuit flow (both circuits equal) [kW].

    Also the stage orders of the two-stage cases swapped (``reversed``), because the order is the
    only difference between cases 3/4 and 5/6.
    """
    swapped: dict[int, k.Order] = {
        3: ((1, 2), ()),
        4: ((), (1, 2)),
        5: ((2, 1), ()),
        6: ((), (2, 1)),
    }
    rows = []
    for flow in flows_kg_per_h:
        si = inputs.with_htm_flow((flow, flow)).to_si()
        row: dict[str, object] = {"htm_flow_kg_per_h": flow}
        for case in (3, 4, 5, 6):
            as_drawn = k.evaluate_case(si, case)
            rev = k.evaluate_case(si, case, order=swapped[case])
            row[f"case{case}_min_stage_kw"] = _min_stage(as_drawn)[0] / 1e3
            row[f"case{case}_reversed_min_stage_kw"] = _min_stage(rev)[0] / 1e3
            row[f"case{case}_gas_kg_per_h"] = as_drawn.total_gas_kg_per_h
            row[f"case{case}_reversed_gas_kg_per_h"] = rev.total_gas_kg_per_h
        rows.append(row)
    return rows


def equal_outlet_vs_counter_current(inputs: ExcelInputsV41) -> list[dict[str, object]]:
    """Per stage: recovered heat, equal outlet temperatures vs the counter-current bound.

    Comparable stage by stage only for the first stage of each sink: a second stage sees the gas
    the first stage delivered, which is warmer in the bound (``sink_in_*`` shows both).

    The bound ignores condensation (``condenses`` flags the stages where the bound's exhaust
    would leave below its dew point, so the bound would have to be higher still).
    """
    si = inputs.to_si()
    rows = []
    for case in range(1, 7):
        eq = k.evaluate_case(si, case)
        cc = k.evaluate_case(si, case, mode="counter_current")
        for number, (stage_eq, stage_cc) in enumerate(zip(eq.stages, cc.stages, strict=True), 1):
            rows.append(
                {
                    "case": case,
                    "stage_of_case": number,
                    "source": stage_eq.source,
                    "sink_in_celsius_equal_outlet": stage_eq.sink_in_celsius,
                    "sink_in_celsius_counter_current": stage_cc.sink_in_celsius,
                    "equal_outlet_kw": stage_eq.recovered_w / 1e3,
                    "counter_current_kw": stage_cc.recovered_w / 1e3,
                    "ratio": stage_cc.recovered_w / stage_eq.recovered_w,
                    "counter_current_exhaust_out_celsius": stage_cc.exhaust_out_celsius,
                    "dew_point_celsius": stage_cc.exhaust_dew_point_celsius,
                    "condenses": stage_cc.condenses,
                }
            )
    return rows


DOCUMENT_TYPOS: list[dict[str, str]] = [
    {
        "item": "D1",
        "where": "equation (2.1)",
        "document": "right-hand side: M_w cp_w T_w,in (twice, with the inlet on both sides)",
        "correct": "right-hand side: M_w cp_w T_w,out",
        "check": "(2.3), (2.4) follow from (2.1) only with T_w,out; the test closes (2.1) with it",
    },
    {
        "item": "D2",
        "where": "equation (3.1)",
        "document": "left: M_g [cp_g T_g,in + Y (dh + cp_v T_g,out)]",
        "correct": "left: M_g [cp_g T_g,in + Y (dh + cp_v T_g,in)]",
        "check": "(3.2) follows from (3.1) only with T_g,in; the test closes (3.1) with it",
    },
    {
        "item": "D3",
        "where": "text under the dryer figure",
        "document": "'the exhaust gas contains moisture of Y_in'",
        "correct": "Y_out (the exhaust moisture of (1.3))",
        "check": "wording only",
    },
    {
        "item": "D4",
        "where": "text before (1.3)",
        "document": "'Combining 3 and 4'",
        "correct": "combining (1.1) and (1.2)",
        "check": "wording only",
    },
    {
        "item": "D5",
        "where": "(1.4), (1.5), (2.1), (2.3): M_gas; elsewhere M_g",
        "document": "two symbols for the dry process gas flow",
        "correct": "one symbol",
        "check": "notation only",
    },
]
"""Typos of the equation document (D1, D2 change an equation; D3-D5 are wording and notation)."""

EQUATION_TO_CELLS: dict[str, str] = {
    "(1.1) dry-solid balance": "Balances E6, E13",
    "(1.3) outlet moisture": "Balances F4, F11 (E7, E14: equilibrium moisture)",
    "(1.5) outlet temperature": "Balances F5, F12 (E8, E15 repeat it)",
    "(2.1)-(2.4) single exchanger, with and without condensation": (
        "not used as such: the sheet uses the two-exchanger form (4.1)/(4.2) without condensation"
    ),
    "(2.2) condensate": "not used (no condensate cell)",
    "(2.5) recovered energy": "Scenarios rows 20, 24, 33, 37 (Mw cp_w (T2 - T1) / 3600 in kW)",
    "(3.2) fuel mass flow": "Scenarios rows 26, 39 (finding G)",
    "(4.1)/(4.2) circuit": (
        "Balances I3:T8 (K1..K6, the coefficients of the 2x2 system) and Scenarios rows 5, 7, 11, "
        "13 (its solution for T2 and T1)"
    ),
}
"""Which cells realise which equation of the document (read from the formulas)."""
