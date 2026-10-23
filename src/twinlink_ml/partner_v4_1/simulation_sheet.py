"""Reader of the sheet ``Simulation``: six cases of the two-line network, Excel against Dyssol.

The sheet (in the partner's results workbook next to the four sheets of ``Comparison.xlsx``) has

* an input block, columns ``A``-``H``: ``Process 1`` / ``Process 2`` in column ``A`` (the line of
  dryer 1 / dryer 2), the component in ``B`` (``Heat recovery``, ``Heating transfer medium``,
  ``Dryer``, ``Heater``), the stream in ``C``, the parameter in ``D``, the unit in ``E``, the value
  in ``F`` (Dyssol's SI unit and value in ``G``, ``H``);
* a result block, columns ``J``-``W``: the quantity in ``J``, an index 1 or 2 in ``K``, and per case
  ``n`` the Excel value in column ``L + 2 (n - 1)`` and the Dyssol value next to it. The medium
  temperatures are stored as text ``"hot - cold"``; the total fuel has ``(kg/h)`` in ``K``.

**What the index ``K`` means is not fixed by the sheet**: the product temperature is per dryer and
the fuel per heater, but the energy and medium rows do not follow one rule over the cases (see
docs/11 section 13). The reader keeps the index as it is; the assignment to a source is made by
the comparison, not here. Rows are found by their labels; a missing label or case raises.

Confidential partner material, never in the repository; the reader takes a path.
"""

from __future__ import annotations

import re
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .comparison import Reported

__all__ = [
    "CASES",
    "RESULT_KEYS",
    "SimulationSheet",
    "load_simulation_sheet",
    "parse_simulation_rows",
]

CASES = (1, 2, 3, 4, 5, 6)
SHEET = "Simulation"

_RESULT_LABELS = (
    ("product stream temperature", "product_celsius"),
    ("exhaust gas temperature", "exhaust_out_celsius"),
    ("recovery", "recovered_kw"),
    ("heat transfer media temperature", "medium"),
    ("total fuel consumption", "total_fuel_kg_per_h"),
    ("fuel consumption", "fuel_kg_per_h"),
)
RESULT_KEYS = (
    "product_celsius",
    "exhaust_out_celsius",
    "recovered_kw",
    "medium_hot_celsius",
    "medium_cold_celsius",
    "fuel_kg_per_h",
)
"""Result keys with an index 1/2; the total fuel is :attr:`SimulationSheet.total_fuel`."""
_PAIR = re.compile(r"^\s*(-?\d+(?:\.\d+)?)\s*-\s*(-?\d+(?:\.\d+)?)\s*$")


@dataclass(frozen=True)
class SimulationSheet:
    """Inputs ``(process, component, stream, parameter) -> value`` and the six cases."""

    inputs: Mapping[tuple[int, str, str, str], float]
    """Sheet units (kg/h, g/kg, degC, fraction, percent, kJ/kg); empty strings for absent labels."""
    results: Mapping[str, Mapping[int, Mapping[int, Reported]]]
    """``key -> case -> index (1, 2) -> (Excel, Dyssol)``."""
    total_fuel: Mapping[int, Reported]
    """``case -> (Excel, Dyssol)`` of the total fuel row (a formula of the two fuel rows)."""

    def value(self, process: int, component: str, stream: str, parameter: str) -> float:
        try:
            return self.inputs[(process, component, stream, parameter)]
        except KeyError:
            raise KeyError(f"no input {process}/{component!r}/{stream!r}/{parameter!r}") from None


def _text(value: object) -> str:
    return value.strip() if isinstance(value, str) else ""


def _number(value: object) -> float:
    if isinstance(value, bool) or not isinstance(value, int | float):
        raise ValueError(f"not a number: {value!r}")
    return float(value)


def _pair(value: object) -> tuple[float, float]:
    match = _PAIR.match(value) if isinstance(value, str) else None
    if match is None:
        raise ValueError(f"not a 'hot - cold' pair: {value!r}")
    return float(match.group(1)), float(match.group(2))


def _result_key(label: str) -> str | None:
    low = label.lower()
    return next((key for text, key in _RESULT_LABELS if text in low), None)


def parse_simulation_rows(rows: list[tuple[Any, ...]]) -> SimulationSheet:
    """Parse the stored values of the sheet (rows as ``openpyxl`` returns them, from ``A1``)."""
    inputs: dict[tuple[int, str, str, str], float] = {}
    results: dict[str, dict[int, dict[int, Reported]]] = {}
    total: dict[int, Reported] = {}
    process = 0
    component = stream = ""
    current: str | None = None
    for row in rows:
        cells = list(row) + [None] * (23 - len(row))
        # ---- input block
        a, b, c, d = (_text(cells[i]) for i in range(4))
        if a.startswith("Process"):
            process, component, stream = int(a.split()[-1]), "", ""
        if b:
            component, stream = b, ""
        if c:
            stream = c
        if process and _text(cells[4]) and isinstance(cells[5], int | float):
            key = (process, component, stream, d)
            if key in inputs:
                raise ValueError(f"input {key} twice")
            inputs[key] = float(cells[5])
        # ---- result block
        label = _text(cells[9])
        if label:
            current = _result_key(label)
            if current is None and cells[10] is not None:
                raise ValueError(f"unknown result label {label!r}")
        if current is None or cells[10] is None:
            continue
        values = {n: (cells[11 + 2 * (n - 1)], cells[12 + 2 * (n - 1)]) for n in CASES}
        if current == "total_fuel_kg_per_h":
            for n, (e, y) in values.items():
                total[n] = Reported(_number(e), _number(y))
            continue
        index = int(_number(cells[10]))
        if current == "medium":
            for n, (e, y) in values.items():
                (eh, ec), (yh, yc) = _pair(e), _pair(y)
                results.setdefault("medium_hot_celsius", {}).setdefault(n, {})[index] = Reported(
                    eh, yh
                )
                results.setdefault("medium_cold_celsius", {}).setdefault(n, {})[index] = Reported(
                    ec, yc
                )
            continue
        for n, (e, y) in values.items():
            results.setdefault(current, {}).setdefault(n, {})[index] = Reported(
                _number(e), _number(y)
            )
    for result_key in RESULT_KEYS:
        per_case = results.get(result_key, {})
        if set(per_case) != set(CASES) or any(set(v) != {1, 2} for v in per_case.values()):
            raise ValueError(f"result {result_key}: cases/indices incomplete")
    if set(total) != set(CASES):
        raise ValueError("total fuel incomplete")
    return SimulationSheet(inputs, results, total)


def load_simulation_sheet(path: str | Path) -> SimulationSheet:
    """Read the sheet ``Simulation`` (needs ``openpyxl``; stored values)."""
    import openpyxl

    book = openpyxl.load_workbook(Path(path), data_only=True)
    if SHEET not in book.sheetnames:
        raise ValueError(f"sheet {SHEET!r} missing")
    return parse_simulation_rows(list(book[SHEET].iter_rows(values_only=True)))
