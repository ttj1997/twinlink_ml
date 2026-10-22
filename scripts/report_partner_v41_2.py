r"""Delivery 4.1_2 of the partner model: workbook diff, five data sets, validators, physical mode.

docs/11 section 12. Inputs (confidential, never copied, never committed) come from
``$TWINLINK_LOCAL_DIR`` (or ``<root>/local``):

* ``Model v4.1.xlsx`` (also ``Model_v4_1.xlsx``): the version 4.1_1 of sections 7 and 11,
* ``partner/Model v4.1_2.xlsx`` (also ``Model_v4.1_2.xlsx``): the version 4.1_2,
* ``partner/Comparison.xlsx``: five data sets on the sheets Dryer, Heater, Heat recovery, Combined,
  each with the partner's Excel and Dyssol value,
* ``partner/Simulation.dflw``: the stationary Dyssol flowsheet of one line (data set 5, Combined).

The artifact holds

* ``workbook_diff`` -- inputs and formula texts 4.1_1 -> 4.1_2, and the reproduction of both
  versions with the switch set of 4.1_1 (all ``partner``) and of 4.1_2 (``gas_formula`` corrected),
* ``data_sets`` -- every result of the four sheets recomputed in the partner mode
  (:class:`~twinlink_ml.flowsheet.line.LineMode`), deviation against Excel and against Dyssol,
* ``dyssol_file`` -- the steady state of ``Simulation.dflw`` and the partner mode at the data set's
  inputs and at the (rounded) inputs stored in the file,
* ``validators`` -- checks (a), (b), (c) of :mod:`twinlink_ml.flowsheet.validators` on the Excel and
  Dyssol values of the tables, on the Dyssol file and on every computed variant,
* ``physical_mode`` -- the same data sets with counter-current exchangers (10 K per exchanger) and
  the saturation limit, medium (i) unpressurised up to 95 degC and (ii) pressurised.

The artifact is never overwritten (new name per run). Run from a clean working tree, started with
``tools/run_here.sh``.
"""

from __future__ import annotations

import argparse
import dataclasses
import json
import math
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

from twinlink_ml.flowsheet import Flowsheet, FlowsheetConfig, Switches
from twinlink_ml.flowsheet.burner import Burner, BurnerConfig, BurnerInputs
from twinlink_ml.flowsheet.dryer import DryerInputs, DryerOut
from twinlink_ml.flowsheet.line import Line, LineInputs, LineMode, LineResult, with_heater
from twinlink_ml.flowsheet.loop import StageOut
from twinlink_ml.flowsheet.sheet import sheet_scenario_cells
from twinlink_ml.flowsheet.thermo import partner_props
from twinlink_ml.flowsheet.validators import (
    Finding,
    SaturationProps,
    adiabatic_saturation_temperature_celsius,
    check_exhaust_above_adiabatic_saturation,
    check_exhaust_saturation,
    check_medium_boiling,
    check_product_mass,
)
from twinlink_ml.partner_v4 import corrected as c
from twinlink_ml.partner_v4.inputs import ATM_PA, SiDryer, SiMaterial
from twinlink_ml.partner_v4_1 import as_built
from twinlink_ml.partner_v4_1 import deviations as dev
from twinlink_ml.partner_v4_1.comparison import (
    DATA_SETS,
    ComparisonSheet,
    DyssolFile,
    load_comparison,
    load_dyssol_steady_state,
    reported_decimals,
)
from twinlink_ml.partner_v4_1.dryer_block import BlockInputs, dryer_block
from twinlink_ml.partner_v4_1.inputs import load_excel_inputs_v41
from twinlink_ml.utils.code_location import require_code_location

ROOT = Path(__file__).resolve().parents[1]
LOCAL = Path(os.environ.get("TWINLINK_LOCAL_DIR", ROOT / "local"))
UNPRESSURISED_MAX_CELSIUS = 95.0
DT_MIN_K = 10.0
MODES: dict[str, LineMode] = {
    "partner": LineMode.partner(),
    "partner_with_saturation_limit": LineMode("partner_with_saturation_limit", "partner", True),
    "physical_unpressurised_95": LineMode.physical(
        htm_max_celsius=UNPRESSURISED_MAX_CELSIUS, dt_min_k=DT_MIN_K
    ),
    "physical_pressurised": LineMode.physical(dt_min_k=DT_MIN_K),
}
SHEET_GAS = {"Dryer": "Process air", "Heater": "Process air", "Combined": "Process air"}


def _git(*args: str) -> str:
    return subprocess.run(
        ["git", *args], cwd=ROOT, capture_output=True, text=True, check=True
    ).stdout


def _find(*names: str) -> Path:
    for name in names:
        if (LOCAL / name).exists():
            return LOCAL / name
    raise FileNotFoundError(f"none of {names} in {LOCAL}")


def _finite(x: float) -> float | None:
    return x if math.isfinite(x) else None


# ---------------------------------------------------------------------- 1. workbook diff
def _row(address: str) -> int:
    return int(re.sub(r"[A-Z]+", "", address))


def _compare(cells: dict[str, float], excel: dict[str, tuple[str, float]]) -> dict[str, Any]:
    if set(cells) != set(excel):
        raise ValueError("cell sets differ")
    rel = {
        a: abs(cells[a] - v) / max(abs(v), 1e-300) if v != 0.0 else abs(cells[a])
        for a, (_, v) in excel.items()
    }
    differing = sorted(
        (a for a in rel if rel[a] > 1e-9 and abs(cells[a] - excel[a][1]) > 1e-12),
        key=lambda a: (_row(a), a),
    )
    agree = [a for a in rel if a not in differing]
    ratios = [excel[a][1] / cells[a] for a in differing if cells[a] != 0.0]
    return {
        "cells": len(excel),
        "largest_relative_error_of_agreeing_cells": max((rel[a] for a in agree), default=0.0),
        "largest_absolute_error_of_agreeing_cells": max(
            (abs(cells[a] - excel[a][1]) for a in agree), default=0.0
        ),
        "differing_cells": differing,
        "differing_rows": sorted({_row(a) for a in differing}),
        "excel_over_computed_in_differing_cells": {
            "min": min(ratios, default=None),
            "max": max(ratios, default=None),
        },
    }


def _formula_diff(old: Path, new: Path) -> list[dict[str, Any]]:
    import openpyxl

    a, b = openpyxl.load_workbook(old), openpyxl.load_workbook(new)
    out = []
    for name in b.sheetnames:
        if name == "Inputs":
            continue
        cells = {c_.coordinate for ws in (a[name], b[name]) for r in ws.iter_rows() for c_ in r}
        for address in sorted(cells, key=lambda x: (_row(x), x)):
            fa, fb = a[name][address].value, b[name][address].value
            if fa != fb:
                out.append({"cell": f"{name}!{address}", "old": fa, "new": fb})
    return out


def workbook_diff(old_path: Path, new_path: Path) -> dict[str, Any]:
    old, new = load_excel_inputs_v41(old_path), load_excel_inputs_v41(new_path)
    switch_sets = {
        "v41_1_all_partner": Switches.all_partner(),
        "v41_2_gas_formula_corrected": Switches(gas_formula="physical"),
    }
    reproduction: dict[str, Any] = {}
    for version, inputs, path in (("v41_1", old, old_path), ("v41_2", new, new_path)):
        excel = dev.read_formula_cells(path, "Scenarios")
        for label, switches in switch_sets.items():
            cells = sheet_scenario_cells(Flowsheet(FlowsheetConfig(inputs.to_si(), switches)))
            reproduction[f"{version}_workbook__{label}_switches"] = _compare(cells, excel)
        balances = as_built.balances_cells_as_built(inputs)
        for sheet in ("Balances", "Balance"):
            stored = dev.read_formula_cells(path, sheet)
            reproduction[f"{version}_workbook__as_built_{sheet}"] = _compare(
                {a: balances[a] for a in stored}, stored
            )
    return {
        "old_workbook": old_path.name,
        "new_workbook": new_path.name,
        "inputs_diff": dev.inputs_diff(old_path, new_path),
        "formula_diff": _formula_diff(old_path, new_path),
        "reference_scan_new": dev.reference_scan(new_path),
        "switch_sets": {k: v.as_dict() for k, v in switch_sets.items()},
        "reproduction": reproduction,
        "inputs_sheet_units_new": {
            "plant": dataclasses.asdict(new.plant),
            "htm_flow_kg_per_h": list(new.htm_flow_kg_per_h),
        },
    }


# ---------------------------------------------------------------------- 2. data sets
def _dryer(sheet: ComparisonSheet, k: int) -> SiDryer:
    gas = SHEET_GAS[sheet.name]
    return SiDryer(
        sheet.value(k, gas, "Mass flow rate") / 3600.0,
        sheet.value(k, gas, "Moisture content") / 1e3,
        sheet.value(k, gas, "Temperature"),
        sheet.value(k, "Spray rate", "Mass flow rate") / 3600.0,
        sheet.value(k, "Spray rate", "Solid fraction"),
        sheet.value(k, "Spray rate", "Temperature"),
    )


def _stream(sheet: ComparisonSheet, k: int, name: str) -> c.Stream:
    return c.Stream(
        sheet.value(k, name, "Mass flow rate") / 3600.0,
        sheet.value(k, name, "Moisture content") / 1e3,
        sheet.value(k, name, "Temperature"),
    )


def _heater_material(sheet: ComparisonSheet, k: int, material: SiMaterial) -> SiMaterial:
    return with_heater(
        material,
        sheet.value(k, "Heater", "Efficiency") / 100.0,
        sheet.value(k, "Heater", "Calorific value") * 1e3,
    )


def _deviation(ours: float, reported: float) -> dict[str, Any]:
    dec = reported_decimals(reported)
    dev_ = ours - reported
    half = 0.5 * 10.0 ** (-dec) if dec is not None else 0.0
    return {
        "reported": reported,
        "deviation": dev_,
        "relative_deviation": dev_ / reported if reported else None,
        "reported_decimals": dec,
        "within_rounding": abs(dev_) <= half + 1e-9 * max(1.0, abs(reported)),
    }


def _quantities(sheet: ComparisonSheet, k: int, ours: dict[str, float]) -> dict[str, Any]:
    out = {}
    for key, per_set in sheet.results.items():
        r = per_set[k]
        out[key] = {
            "ours": ours[key],
            "excel": _deviation(ours[key], r.excel),
            "dyssol": _deviation(ours[key], r.dyssol),
            "excel_minus_dyssol": r.excel - r.dyssol,
        }
    return out


class Calculator:
    """The four sheets in one :class:`LineMode`."""

    def __init__(
        self, mode: LineMode, material: SiMaterial, cp_htm: float, ambient_c: float
    ) -> None:
        self.mode, self.material, self.cp_htm, self.ambient_c = mode, material, cp_htm, ambient_c
        self.line = Line(mode)

    def dryer(self, sheet: ComparisonSheet, k: int) -> DryerOut:
        inputs = LineInputs(_dryer(sheet, k), self.ambient_c, self.material, 0.0, self.cp_htm)
        return self.line.dryer_unit(inputs).solve(DryerInputs(inputs.dryer))

    def heater(self, sheet: ComparisonSheet, k: int) -> float:
        m = _heater_material(sheet, k, self.material)
        gas = c.Stream(
            sheet.value(k, "Process air", "Mass flow rate") / 3600.0,
            sheet.value(k, "Process air", "Moisture content") / 1e3,
            sheet.value(k, "Process air", "Inlet temperature"),
        )
        burner = Burner(BurnerConfig(m, partner_props(m), "physical"))
        out = burner.solve(BurnerInputs(gas, sheet.value(k, "Process air", "Exit temperature")))
        return out.fuel_corrected_kg_per_s * 3600.0

    def recovery(self, sheet: ComparisonSheet, k: int) -> StageOut:
        return self.line.recovery(
            self.material,
            _stream(sheet, k, "Dryer exhaust"),
            _stream(sheet, k, "Fresh gas"),
            sheet.value(k, "Heating transfer medium", "Flow rate") / 3600.0,
            self.cp_htm,
        )

    def combined(self, sheet: ComparisonSheet, k: int) -> LineResult:
        dryer = _dryer(sheet, k)
        fresh = _stream(sheet, k, "Fresh gas")
        if (fresh.dry_air_kg_per_s, fresh.humidity_kg_per_kg) != (
            dryer.air_mass_flow_kg_per_s,
            dryer.air_humidity_kg_per_kg,
        ):
            raise ValueError(f"Combined data set {k}: fresh gas and process air differ")
        inputs = LineInputs(
            dryer,
            fresh.temperature_celsius,
            _heater_material(sheet, k, self.material),
            sheet.value(k, "Heating transfer medium", "Flow rate") / 3600.0,
            self.cp_htm,
        )
        return self.line.solve(inputs)


def _stage_dict(s: StageOut) -> dict[str, Any]:
    return {
        "recovered_kw": s.q_w / 1e3,
        "gas_in_celsius": s.gas_in_celsius,
        "gas_out_celsius": s.gas_out_celsius,
        "exhaust_in_celsius": s.exhaust_in_celsius,
        "exhaust_out_celsius": s.exhaust_out_celsius,
        "medium_cold_celsius": s.htm_cold_celsius,
        "medium_hot_celsius": s.htm_hot_celsius,
        "binding": s.binding,
        "min_approach_recovery_k": _finite(s.min_approach_recovery_k),
        "min_approach_preheater_k": s.min_approach_preheater_k,
        "exhaust_dew_point_celsius": s.exhaust_dew_point_celsius,
        "condensate_modelled_kg_per_h": s.condensate_modelled_kg_per_h,
        "condensate_unmodelled_kg_per_h": s.condensate_unmodelled_kg_per_h,
    }


def _dryer_dict(d: DryerOut) -> dict[str, Any]:
    return {
        "exhaust_celsius": d.exhaust.temperature_celsius,
        "exhaust_humidity_g_per_kg": d.exhaust.humidity_kg_per_kg * 1e3,
        "evaporation_kg_per_h": d.evaporation_kg_per_s * 3600.0,
        "product_kg_per_h": d.product_kg_per_s * 3600.0,
        "product_moisture_g_per_kg": d.product_moisture_kg_per_kg * 1e3,
        "saturation_limited": d.saturation_limited,
    }


def _line_dict(r: LineResult) -> dict[str, Any]:
    return {
        "dryer": _dryer_dict(r.dryer),
        "stage": _stage_dict(r.stage),
        "recovered_kw": r.recovered_kw,
        "heater_in_celsius": r.stage.gas_out_celsius,
        "heater_duty_kw": r.heater_duty_kw,
        "fuel_kg_per_h": r.fuel_kg_per_h,
        "product_kg_per_h": r.product_kg_per_h,
        "product_celsius": r.product_celsius,
    }


def _findings(items: list[Finding]) -> list[dict[str, object]]:
    return [f.as_dict() for f in items]


def _dryer_checks(
    exhaust_c: float, y_out: float, product_kg_per_s: float, d: SiDryer, sat: SaturationProps
) -> list[Finding]:
    return [
        check_exhaust_saturation(exhaust_c, y_out),
        check_exhaust_above_adiabatic_saturation(
            exhaust_c, d.air_temperature_celsius, d.air_humidity_kg_per_kg, sat
        ),
        check_product_mass(
            product_kg_per_s, d.spray_mass_flow_kg_per_s, d.spray_solid_fraction_kg_per_kg
        ),
    ]


def _reported_dryer_checks(
    exhaust_c: float, product_kg_per_h: float, d: SiDryer, sat: SaturationProps
) -> dict[str, Any]:
    """Checks on reported values; the exhaust humidity follows from the reported product mass."""
    evaporation = d.spray_mass_flow_kg_per_s - product_kg_per_h / 3600.0
    y_out = d.air_humidity_kg_per_kg + evaporation / d.air_mass_flow_kg_per_s
    return {
        "exhaust_humidity_from_product_mass_g_per_kg": y_out * 1e3,
        "findings": _findings(_dryer_checks(exhaust_c, y_out, product_kg_per_h / 3600.0, d, sat)),
    }


def _reported_balance(
    sheet: ComparisonSheet, k: int, material: SiMaterial, cp_htm: float, fuel_target_c: float | None
) -> dict[str, Any]:
    """Check whether the reported values close the balances of the equal-outlet loop.

    Recovered heat from the gas side ``C_gas (T_cold - T_fresh)`` and from the medium
    ``B (T_hot - T_cold)``, with the reported temperatures; for Combined also the fuel the
    reported cold side implies. Rounded temperatures (1 K) limit the closure to about
    ``C_gas * 0.5 K`` (about 0.02 kW here).
    """
    fresh = _stream(sheet, k, "Fresh gas")
    props = partner_props(material)
    cg = fresh.dry_air_kg_per_s * (props.cp_air + fresh.humidity_kg_per_kg * props.cp_vapour)
    b = sheet.value(k, "Heating transfer medium", "Flow rate") / 3600.0 * cp_htm
    out = {}
    for src in ("excel", "dyssol"):
        q = getattr(sheet.results["recovered_kw"][k], src)
        hot = getattr(sheet.results["medium_hot_celsius"][k], src)
        cold = getattr(sheet.results["medium_cold_celsius"][k], src)
        entry: dict[str, Any] = {
            "recovered_kw_reported": q,
            "recovered_kw_from_gas_side": cg * (cold - fresh.temperature_celsius) / 1e3,
            "recovered_kw_from_medium": b * (hot - cold) / 1e3,
            "half_kelvin_on_gas_side_kw": cg * 0.5 / 1e3,
        }
        if fuel_target_c is not None:
            m = _heater_material(sheet, k, material)
            entry["fuel_kg_per_h_from_cold_side"] = (
                cg * (fuel_target_c - cold) / (m.heater_efficiency * m.lower_heating_value_j_per_kg)
            ) * 3600.0
            entry["fuel_kg_per_h_reported"] = getattr(sheet.results["fuel_kg_per_h"][k], src)
        out[src] = entry
    return out


def dryer_3_spray_at_adiabatic_saturation(
    sheets: dict[str, ComparisonSheet], material: SiMaterial, cp_htm: float, ambient_c: float
) -> dict[str, Any]:
    """Dryer data set 3 with the saturation limit, spray as given and spray at ``T_as``.

    Check (a2) is strict for an adiabatic dryer whose liquid enters at ``T_as``; with the spray
    at ``T_as`` the saturation-limited block must land on ``T_as`` itself.
    """
    sheet = sheets["Dryer"]
    sat = SaturationProps.from_material(material)
    d = _dryer(sheet, 3)
    t_as = adiabatic_saturation_temperature_celsius(
        d.air_temperature_celsius, d.air_humidity_kg_per_kg, sat
    )
    line = Line(MODES["partner_with_saturation_limit"])
    out: dict[str, Any] = {"adiabatic_saturation_celsius": t_as}
    for label, t_s in (("spray_as_given", d.spray_temperature_celsius), ("spray_at_t_as", t_as)):
        dd = dataclasses.replace(d, spray_temperature_celsius=t_s)
        o = line.dryer_unit(LineInputs(dd, ambient_c, material, 0.0, cp_htm)).solve(DryerInputs(dd))
        out[label] = {
            "spray_celsius": t_s,
            "exhaust_celsius": o.exhaust.temperature_celsius,
            "exhaust_minus_t_as_k": o.exhaust.temperature_celsius - t_as,
        }
    return out


def data_sets_and_checks(
    sheets: dict[str, ComparisonSheet], material: SiMaterial, cp_htm: float, ambient_c: float
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    sat = SaturationProps.from_material(material)
    calc = {name: Calculator(mode, material, cp_htm, ambient_c) for name, mode in MODES.items()}
    partner = calc["partner"]
    data: dict[str, Any] = {}
    checks: dict[str, Any] = {}
    physical: dict[str, Any] = {}
    dr, he, hr, co = (sheets[s] for s in ("Dryer", "Heater", "Heat recovery", "Combined"))
    for k in DATA_SETS:
        ds = str(k)
        # Dryer
        out = partner.dryer(dr, k)
        data.setdefault("Dryer", {})[ds] = {
            "computed": _dryer_dict(out),
            "quantities": _quantities(
                dr,
                k,
                {
                    "product_kg_per_h": out.product_kg_per_s * 3600.0,
                    "exhaust_celsius": out.exhaust.temperature_celsius,
                },
            ),
        }
        d = _dryer(dr, k)
        checks.setdefault("Dryer", {})[ds] = {
            src: _reported_dryer_checks(
                getattr(dr.results["exhaust_celsius"][k], src),
                getattr(dr.results["product_kg_per_h"][k], src),
                d,
                sat,
            )
            for src in ("excel", "dyssol")
        }
        physical.setdefault("Dryer", {})[ds] = {}
        for name, cal in calc.items():
            o = cal.dryer(dr, k)
            physical["Dryer"][ds][name] = {
                **_dryer_dict(o),
                "findings": _findings(
                    _dryer_checks(
                        o.exhaust.temperature_celsius,
                        o.exhaust.humidity_kg_per_kg,
                        o.product_kg_per_s,
                        d,
                        sat,
                    )
                ),
            }
            checks["Dryer"][ds][f"computed_{name}"] = physical["Dryer"][ds][name]["findings"]
        # Heater (no assumption of the modes acts on it)
        fuel = partner.heater(he, k)
        data.setdefault("Heater", {})[ds] = {
            "computed": {"fuel_kg_per_h": fuel},
            "quantities": _quantities(he, k, {"fuel_kg_per_h": fuel}),
        }
        # Heat recovery
        st = partner.recovery(hr, k)
        data.setdefault("Heat recovery", {})[ds] = {
            "computed": _stage_dict(st),
            "quantities": _quantities(
                hr,
                k,
                {
                    "recovered_kw": st.q_w / 1e3,
                    "medium_hot_celsius": st.htm_hot_celsius or math.nan,
                    "medium_cold_celsius": st.htm_cold_celsius or math.nan,
                },
            ),
            "reported_internal_balance": _reported_balance(hr, k, material, cp_htm, None),
        }
        ex = _stream(hr, k, "Dryer exhaust")
        checks.setdefault("Heat recovery", {})[ds] = {
            "given_exhaust": _findings(
                [check_exhaust_saturation(ex.temperature_celsius, ex.humidity_kg_per_kg)]
            ),
            **{
                src: _findings(
                    [check_medium_boiling(getattr(hr.results["medium_hot_celsius"][k], src))]
                )
                for src in ("excel", "dyssol")
            },
        }
        physical.setdefault("Heat recovery", {})[ds] = {}
        for name, cal in calc.items():
            s = cal.recovery(hr, k)
            physical["Heat recovery"][ds][name] = {
                **_stage_dict(s),
                "findings": _findings([check_medium_boiling(s.htm_hot_celsius or -math.inf)]),
            }
            checks["Heat recovery"][ds][f"computed_{name}"] = physical["Heat recovery"][ds][name][
                "findings"
            ]
        # Combined
        line = partner.combined(co, k)
        data.setdefault("Combined", {})[ds] = {
            "computed": _line_dict(line),
            "quantities": _quantities(
                co,
                k,
                {
                    "recovered_kw": line.recovered_kw,
                    "medium_hot_celsius": line.medium_hot_celsius,
                    "medium_cold_celsius": line.medium_cold_celsius,
                    "fuel_kg_per_h": line.fuel_kg_per_h,
                    "product_kg_per_h": line.product_kg_per_h,
                    "product_celsius": line.product_celsius,
                },
            ),
            "reported_internal_balance": _reported_balance(
                co, k, material, cp_htm, _dryer(co, k).air_temperature_celsius
            ),
        }
        dc = _dryer(co, k)
        checks.setdefault("Combined", {})[ds] = {
            src: {
                **_reported_dryer_checks(
                    getattr(co.results["product_celsius"][k], src),
                    getattr(co.results["product_kg_per_h"][k], src),
                    dc,
                    sat,
                ),
                "medium": _findings(
                    [check_medium_boiling(getattr(co.results["medium_hot_celsius"][k], src))]
                ),
            }
            for src in ("excel", "dyssol")
        }
        physical.setdefault("Combined", {})[ds] = {}
        for name, cal in calc.items():
            r = cal.combined(co, k)
            physical["Combined"][ds][name] = {
                **_line_dict(r),
                "findings": _findings(list(r.findings())),
            }
            checks["Combined"][ds][f"computed_{name}"] = physical["Combined"][ds][name]["findings"]
    return data, checks, physical


# ---------------------------------------------------------------------- 3. Dyssol file
_EQUIVALENT_RANGES = {
    "evaporation_enthalpy_0c_j_per_kg": (2.0e6, 2.6e6),
    "cp_vapour_j_per_kg_k": (1000.0, 3000.0),
    "cp_gas_j_per_kg_k": (900.0, 1200.0),
    "cp_liquid_j_per_kg_k": (1000.0, 8000.0),
    "cp_solid_j_per_kg_k": (-5000.0, 5000.0),
}


def _dryer_equivalents(inputs: LineInputs, target_celsius: float) -> dict[str, Any]:
    """The value of one property that alone would give Dyssol's exhaust (an equivalent, no cause)."""
    props = Line(MODES["partner"]).dryer_unit(inputs).properties()
    d = inputs.dryer
    block = BlockInputs(
        d.air_mass_flow_kg_per_s,
        d.air_humidity_kg_per_kg,
        d.air_temperature_celsius,
        d.spray_mass_flow_kg_per_s,
        d.spray_solid_fraction_kg_per_kg,
        d.spray_temperature_celsius,
        inputs.material.equilibrium_moisture_kg_per_kg,
        inputs.fresh_gas_celsius,
    )

    def t_out(name: str, x: float) -> float:
        p = dataclasses.replace(props, **{name: x})
        return dryer_block(block, p).outlet_temperature_celsius

    out: dict[str, Any] = {"target_celsius": target_celsius}
    for name, (lo, hi) in _EQUIVALENT_RANGES.items():
        flo, fhi = t_out(name, lo) - target_celsius, t_out(name, hi) - target_celsius
        if flo * fhi > 0.0:
            out[name] = None
            continue
        for _ in range(200):
            mid = 0.5 * (lo + hi)
            fm = t_out(name, mid) - target_celsius
            lo, hi, flo = (mid, hi, fm) if fm * flo > 0.0 else (lo, mid, flo)
        out[name] = {"value": 0.5 * (lo + hi), "partner_value": getattr(props, name)}
    return out


def dyssol_file(
    f: DyssolFile, sheets: dict[str, ComparisonSheet], material: SiMaterial, cp_htm: float,
    ambient_c: float,
) -> dict[str, Any]:  # fmt: skip
    co = sheets["Combined"]
    heater = f.units["Heater TwinLink"]
    recovery = f.units["Heat recovery TwinLink"]
    dryer_unit = f.units["DryerTwinLink"]
    gas_dry, gas_y = f.gas_split("Gas In")
    out_dry, out_y = f.gas_split("Dryer Out")
    feed = f.streams["Feed"]
    product = f.streams["Product"]
    eff = heater.parameters["Efficiency"][0] / 100.0
    cv = heater.parameters["Calorific Value"][0] * 1e3
    from_file = LineInputs(
        SiDryer(
            gas_dry,
            gas_y,
            heater.parameters["Outlet Temperature"][0],
            feed.mass_flow_kg_per_s,
            feed.phase_fractions["Solid"],
            feed.temperature_celsius,
        ),
        f.streams["Gas In"].temperature_celsius,
        dataclasses.replace(
            with_heater(material, eff, cv),
            equilibrium_moisture_kg_per_kg=dryer_unit.parameters["Equilibrium moisture content"][0]
            / 1e3,
        ),
        recovery.parameters["Medium mass flow"][0] / 3600.0,
        cp_htm,
    )
    dyssol = {
        "recovered_kw": recovery.state["Heat recovered [kW]"],
        "medium_hot_celsius": recovery.state["Medium hot side temperature [C]"],
        "medium_cold_celsius": recovery.state["Medium cold side temperature [C]"],
        "heater_in_celsius": f.streams["Heater In"].temperature_celsius,
        "heater_duty_kw": heater.state["Heat duty [kW]"],
        "fuel_kg_per_h": heater.state["Fuel consumption [kg/h]"],
        "product_kg_per_h": product.mass_flow_kg_per_s * 3600.0,
        "product_celsius": product.temperature_celsius,
        "exhaust_celsius": f.streams["Dryer Out"].temperature_celsius,
        "exhaust_humidity_g_per_kg": out_y * 1e3,
        "exhaust_dry_gas_kg_per_h": out_dry * 3600.0,
        "exhaust_after_recovery_celsius": f.streams["Gas Out"].temperature_celsius,
    }
    partner = Calculator(MODES["partner"], material, cp_htm, ambient_c)

    def flat(r: LineResult) -> dict[str, float]:
        return {
            "recovered_kw": r.recovered_kw,
            "medium_hot_celsius": r.medium_hot_celsius,
            "medium_cold_celsius": r.medium_cold_celsius,
            "heater_in_celsius": r.stage.gas_out_celsius,
            "heater_duty_kw": r.heater_duty_kw,
            "fuel_kg_per_h": r.fuel_kg_per_h,
            "product_kg_per_h": r.product_kg_per_h,
            "product_celsius": r.product_celsius,
            "exhaust_celsius": r.dryer.exhaust.temperature_celsius,
            "exhaust_humidity_g_per_kg": r.dryer.exhaust.humidity_kg_per_kg * 1e3,
            "exhaust_dry_gas_kg_per_h": r.inputs.dryer.air_mass_flow_kg_per_s * 3600.0,
            "exhaust_after_recovery_celsius": r.stage.exhaust_out_celsius,
        }

    runs = {
        "data_set_inputs": flat(partner.combined(co, 5)),
        "file_inputs": flat(Line(MODES["partner"]).solve(from_file)),
    }
    # the units one by one at Dyssol's own states: which unit carries the difference
    at_exhaust = Line(MODES["partner"]).recovery(
        from_file.material,
        c.Stream(out_dry, out_y, dyssol["exhaust_celsius"]),
        from_file.fresh_gas,
        from_file.htm_kg_per_s,
        cp_htm,
    )
    heater_in = c.Stream(gas_dry, gas_y, dyssol["heater_in_celsius"])
    burner = (
        Line(MODES["partner"])
        .burner_unit(from_file.material)
        .solve(BurnerInputs(heater_in, from_file.dryer.air_temperature_celsius))
    )
    unit_by_unit = {
        "recovery_at_dyssol_exhaust": {
            **_stage_dict(at_exhaust),
            "minus_dyssol_kw": at_exhaust.q_w / 1e3 - dyssol["recovered_kw"],
        },
        "heater_at_dyssol_heater_inlet": {
            "duty_kw": burner.duty_w / 1e3,
            "fuel_kg_per_h": burner.fuel_corrected_kg_per_s * 3600.0,
            "duty_minus_dyssol_kw": burner.duty_w / 1e3 - dyssol["heater_duty_kw"],
        },
        "dryer_single_parameter_equivalents": _dryer_equivalents(
            from_file, dyssol["exhaust_celsius"]
        ),
    }
    gas_cap = gas_dry * (material.cp_gas_j_per_kg_k + gas_y * material.cp_vapour_j_per_kg_k)
    ex_cap = out_dry * (material.cp_gas_j_per_kg_k + out_y * material.cp_vapour_j_per_kg_k)
    q = dyssol["recovered_kw"] * 1e3
    balances = {
        "gas_side_kw_with_partner_cp": gas_cap
        * (dyssol["heater_in_celsius"] - f.streams["Gas In"].temperature_celsius) / 1e3,
        "exhaust_side_kw_with_partner_cp": ex_cap
        * (dyssol["exhaust_celsius"] - dyssol["exhaust_after_recovery_celsius"]) / 1e3,
        "medium_kw_with_cp_4_185": recovery.parameters["Medium mass flow"][0] / 3600.0 * cp_htm
        * (dyssol["medium_hot_celsius"] - dyssol["medium_cold_celsius"]) / 1e3,
        "heater_kw_with_partner_cp": gas_cap
        * (heater.parameters["Outlet Temperature"][0] - dyssol["heater_in_celsius"]) / 1e3,
        "fuel_from_duty_kg_per_h": dyssol["heater_duty_kw"] * 1e3 / (eff * cv) * 3600.0,
        "recovered_kw": q / 1e3,
    }  # fmt: skip
    p = ATM_PA
    sat = SaturationProps.from_material(material)
    d = from_file.dryer
    return {
        "compounds": list(f.compounds),
        "vapour_compound_index": f.vapour_index(),
        "streams": {k: dataclasses.asdict(v) for k, v in f.streams.items()},
        "units": {k: dataclasses.asdict(v) for k, v in f.units.items()},
        "file_inputs": dataclasses.asdict(from_file),
        "dyssol": dyssol,
        "computed": runs,
        "deviation_computed_minus_dyssol": {
            run: {key: values[key] - dyssol[key] for key in dyssol} for run, values in runs.items()
        },
        "dyssol_internal_balances": balances,
        "unit_by_unit": unit_by_unit,
        "findings": _findings(
            [
                *_dryer_checks(
                    dyssol["exhaust_celsius"], out_y, product.mass_flow_kg_per_s, d, sat
                ),
                check_medium_boiling(dyssol["medium_hot_celsius"], pressure_pa=p),
            ]
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--name", default="partner_v41_2_comparison")
    parser.add_argument("--out-dir", type=Path, default=ROOT / "artifacts" / "runs")
    args = parser.parse_args()
    target = args.out_dir / f"{args.name}.json"
    if target.exists():
        print(f"{target} exists; choose a new --name (artifacts are never overwritten)")
        return 1
    old_path = _find("Model v4.1.xlsx", "Model_v4_1.xlsx")
    new_path = _find("partner/Model v4.1_2.xlsx", "partner/Model_v4.1_2.xlsx")
    comparison_path = _find("partner/Comparison.xlsx")
    dflw_path = _find("partner/Simulation.dflw")
    new = load_excel_inputs_v41(new_path)
    si = new.to_si()
    material, cp_htm = si.plant.material, si.cp_htm_j_per_kg_k
    ambient_c = si.plant.ambient.temperature_celsius
    sheets = load_comparison(comparison_path)
    data, checks, physical = data_sets_and_checks(sheets, material, cp_htm, ambient_c)
    record: dict[str, Any] = {
        "provenance": {
            "git/commit": _git("rev-parse", "HEAD").strip(),
            "git/dirty": bool(_git("status", "--porcelain").strip()),
            "python": sys.version.split()[0],
            **require_code_location(ROOT),
        },
        "parameters": {
            "files": [p.name for p in (old_path, new_path, comparison_path, dflw_path)],
            "material_from": new_path.name,
            "material_si": dataclasses.asdict(material),
            "cp_htm_j_per_kg_k": cp_htm,
            "dryer_sheet_ambient_celsius": ambient_c,
            "modes": {k: dataclasses.asdict(v) for k, v in MODES.items()},
            "dt_min_convention": "per exchanger",
            "saturation": "IAPWS-IF97 region 4, 1 atm, Y_sat = 0.62198 p_s / (p - p_s)",
            "within_rounding": "|ours - reported| <= half a unit of the reported last decimal",
        },
        "comparison_inputs_sheet_units": {
            name: {str(k): {f"{a}/{b}": v for (a, b), v in s.inputs[k].items()} for k in DATA_SETS}
            for name, s in sheets.items()
        },
        "workbook_diff": workbook_diff(old_path, new_path),
        "data_sets": data,
        "dyssol_file": dyssol_file(
            load_dyssol_steady_state(dflw_path), sheets, material, cp_htm, ambient_c
        ),
        "validators": checks,
        "dryer_3_spray_at_adiabatic_saturation": dryer_3_spray_at_adiabatic_saturation(
            sheets, material, cp_htm, ambient_c
        ),
        "physical_mode": physical,
    }
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(record, indent=1, default=str), encoding="utf-8")
    print(f"wrote {target}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
