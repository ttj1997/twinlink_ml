"""Delivery 4.1_2 of the partner model: reproduction, the five data sets, Dyssol, validators.

Needs the confidential files from ``local/`` (``TWINLINK_LOCAL_DIR`` names the directory, e.g. in
a worktree; never copied, never committed): ``Model v4.1.xlsx`` (4.1_1),
``partner/Model v4.1_2.xlsx``,
``partner/Comparison.xlsx``, ``partner/Simulation.dflw``. Each test is skipped when its file is
missing. Input and expected values are read from those files only; the reader of the comparison
workbook is tested on a synthetic workbook as well.
"""

from __future__ import annotations

import math
import os
import re
from pathlib import Path

import pytest

from twinlink_ml.flowsheet import Flowsheet, FlowsheetConfig, Switches
from twinlink_ml.flowsheet.burner import BurnerInputs
from twinlink_ml.flowsheet.dryer import DryerInputs
from twinlink_ml.flowsheet.line import Line, LineInputs, LineMode, with_heater
from twinlink_ml.flowsheet.sheet import sheet_scenario_cells
from twinlink_ml.flowsheet.validators import (
    SaturationProps,
    check_exhaust_saturation,
    check_medium_boiling,
    check_product_mass,
)
from twinlink_ml.partner_v4 import corrected as c
from twinlink_ml.partner_v4.inputs import SiDryer
from twinlink_ml.partner_v4_1 import as_built
from twinlink_ml.partner_v4_1 import deviations as dev
from twinlink_ml.partner_v4_1.comparison import (
    DATA_SETS,
    ComparisonSheet,
    load_comparison,
    load_dyssol_steady_state,
    reported_decimals,
)
from twinlink_ml.partner_v4_1.inputs import load_excel_inputs_v41

LOCAL = Path(os.environ.get("TWINLINK_LOCAL_DIR", Path(__file__).resolve().parents[1] / "local"))


def _first(*names: str) -> Path | None:
    return next((LOCAL / n for n in names if (LOCAL / n).exists()), None)


OLD = _first("Model v4.1.xlsx", "Model_v4_1.xlsx")
NEW = _first("partner/Model v4.1_2.xlsx", "partner/Model_v4.1_2.xlsx")
COMPARISON = _first("partner/Comparison.xlsx")
DFLW = _first("partner/Simulation.dflw")


def _need(path: Path | None, module: str = "openpyxl") -> Path:
    if path is None:
        pytest.skip("partner file of the delivery 4.1_2 not available in local/")
    pytest.importorskip(module)
    return path


@pytest.fixture(scope="module")
def new_inputs():  # type: ignore[no-untyped-def]
    return load_excel_inputs_v41(_need(NEW))


@pytest.fixture(scope="module")
def sheets() -> dict[str, ComparisonSheet]:
    return load_comparison(_need(COMPARISON))


# ---------------------------------------------------------------------- 4.1_1 -> 4.1_2
def test_new_switch_set_reproduces_every_stored_number_of_4_1_2(new_inputs) -> None:  # type: ignore[no-untyped-def]
    cells = sheet_scenario_cells(
        Flowsheet(FlowsheetConfig(new_inputs.to_si(), Switches(gas_formula="physical")))
    )
    excel = dev.read_formula_cells(_need(NEW), "Scenarios")
    assert set(cells) == set(excel)
    for address, (_, stored) in excel.items():
        assert cells[address] == pytest.approx(stored, rel=1e-12, abs=1e-12), address


def test_old_switch_set_on_4_1_2_differs_only_in_the_gas_rows(new_inputs) -> None:  # type: ignore[no-untyped-def]
    cells = sheet_scenario_cells(Flowsheet(FlowsheetConfig(new_inputs.to_si(), Switches())))
    excel = dev.read_formula_cells(_need(NEW), "Scenarios")
    differing = {
        a for a, (_, v) in excel.items() if cells[a] != pytest.approx(v, rel=1e-9, abs=1e-12)
    }
    assert {int(re.sub(r"[A-Z]+", "", a)) for a in differing} == {26, 39, 41}


def test_new_switch_set_on_4_1_1_differs_only_in_the_gas_rows_by_108_91() -> None:
    inputs = load_excel_inputs_v41(_need(OLD))
    cells = sheet_scenario_cells(
        Flowsheet(FlowsheetConfig(inputs.to_si(), Switches(gas_formula="physical")))
    )
    excel = dev.read_formula_cells(OLD, "Scenarios")  # type: ignore[arg-type]
    differing = {
        a for a, (_, v) in excel.items() if cells[a] != pytest.approx(v, rel=1e-9, abs=1e-12)
    }
    assert {int(re.sub(r"[A-Z]+", "", a)) for a in differing} == {26, 39, 41}
    for a in differing:
        assert excel[a][1] / cells[a] == pytest.approx(13.2 / 1.212 * 10.0, rel=1e-9)


@pytest.mark.parametrize("sheet", ["Balances", "Balance"])
def test_as_built_balances_reproduce_4_1_2(new_inputs, sheet: str) -> None:  # type: ignore[no-untyped-def]
    cells = as_built.balances_cells_as_built(new_inputs)
    for address, (_, stored) in dev.read_formula_cells(_need(NEW), sheet).items():
        assert cells[address] == pytest.approx(stored, rel=1e-12, abs=1e-12), address


def test_only_the_gas_formulas_changed_and_column_k_and_balance_are_unused() -> None:
    import openpyxl

    a, b = openpyxl.load_workbook(_need(OLD)), openpyxl.load_workbook(_need(NEW))
    changed = set()
    for name in ("Scenarios", "Balances", "Balance"):
        for row in b[name].iter_rows():
            for cell in row:
                if a[name][cell.coordinate].value != cell.value:
                    changed.add(f"{name}!{cell.coordinate}")
    gas = {f"Scenarios!{col}{row}" for col in "EGIKMO" for row in (26, 39)}
    assert changed == gas | {"Scenarios!E41"}
    for cell in gas:
        formula = b["Scenarios"][cell.split("!")[1]].value
        assert "Inputs!I9/1000" in formula and "Inputs!I6/100)" in formula
    scan = dev.reference_scan(NEW)  # type: ignore[arg-type]
    assert scan["inputs_column_k_referenced_by"] == []
    assert scan["balance_sheet_referenced_by"] == []


# ---------------------------------------------------------------------- the five data sets
def _dryer(sheet: ComparisonSheet, k: int) -> SiDryer:
    return SiDryer(
        sheet.value(k, "Process air", "Mass flow rate") / 3600,
        sheet.value(k, "Process air", "Moisture content") / 1e3,
        sheet.value(k, "Process air", "Temperature"),
        sheet.value(k, "Spray rate", "Mass flow rate") / 3600,
        sheet.value(k, "Spray rate", "Solid fraction"),
        sheet.value(k, "Spray rate", "Temperature"),
    )


def _within(ours: float, reported: float) -> bool:
    dec = reported_decimals(reported)
    half = 0.5 * 10.0 ** (-dec) if dec is not None else 0.0
    return abs(ours - reported) <= half + 1e-9 * max(1.0, abs(reported))


def _combined(sheets, new_inputs, k: int, mode: LineMode):  # type: ignore[no-untyped-def]
    co = sheets["Combined"]
    si = new_inputs.to_si()
    material = with_heater(
        si.plant.material,
        co.value(k, "Heater", "Efficiency") / 100,
        co.value(k, "Heater", "Calorific value") * 1e3,
    )
    inputs = LineInputs(
        _dryer(co, k),
        co.value(k, "Fresh gas", "Temperature"),
        material,
        co.value(k, "Heating transfer medium", "Flow rate") / 3600,
        si.cp_htm_j_per_kg_k,
    )
    return Line(mode).solve(inputs)


def test_dryer_data_set_1_reproduces_the_unrounded_excel_exhaust(sheets, new_inputs) -> None:  # type: ignore[no-untyped-def]
    dr = sheets["Dryer"]
    reported = dr.results["exhaust_celsius"][1].excel
    assert reported_decimals(reported) is None
    si = new_inputs.to_si()
    inputs = LineInputs(_dryer(dr, 1), 20.0, si.plant.material, 0.0, si.cp_htm_j_per_kg_k)
    line = Line(LineMode.partner())
    out = line.dryer_unit(inputs).solve(DryerInputs(inputs.dryer))
    assert out.exhaust.temperature_celsius == pytest.approx(reported, rel=1e-12)


NOT_WITHIN_ROUNDING_OF_EXCEL = {
    ("Combined", 3, "recovered_kw"),
    ("Combined", 3, "medium_hot_celsius"),
    ("Combined", 3, "medium_cold_celsius"),
    ("Combined", 3, "fuel_kg_per_h"),
}
"""Excel's Combined data set 3 does not close its own loop balance (docs/11, section 12)."""


def test_combined_partner_mode_within_rounding_of_excel(sheets, new_inputs) -> None:  # type: ignore[no-untyped-def]
    co = sheets["Combined"]
    for k in DATA_SETS:
        r = _combined(sheets, new_inputs, k, LineMode.partner())
        ours = {
            "recovered_kw": r.recovered_kw,
            "medium_hot_celsius": r.medium_hot_celsius,
            "medium_cold_celsius": r.medium_cold_celsius,
            "fuel_kg_per_h": r.fuel_kg_per_h,
            "product_kg_per_h": r.product_kg_per_h,
            "product_celsius": r.product_celsius,
        }
        for key, value in ours.items():
            expected = ("Combined", k, key) not in NOT_WITHIN_ROUNDING_OF_EXCEL
            assert _within(value, co.results[key][k].excel) == expected, (k, key)


def test_dyssol_recovery_unit_is_the_equal_outlet_loop(new_inputs) -> None:  # type: ignore[no-untyped-def]
    f = load_dyssol_steady_state(_need(DFLW, "h5py"))
    rec = f.units["Heat recovery TwinLink"]
    gas_dry, gas_y = f.gas_split("Gas In")
    ex_dry, ex_y = f.gas_split("Dryer Out")
    si = new_inputs.to_si()
    stage = Line(LineMode.partner()).recovery(
        si.plant.material,
        c.Stream(ex_dry, ex_y, f.streams["Dryer Out"].temperature_celsius),
        c.Stream(gas_dry, gas_y, f.streams["Gas In"].temperature_celsius),
        rec.parameters["Medium mass flow"][0] / 3600,
        si.cp_htm_j_per_kg_k,
    )
    assert stage.q_w / 1e3 == pytest.approx(rec.state["Heat recovered [kW]"], rel=1e-9)
    assert stage.htm_hot_celsius == pytest.approx(
        rec.state["Medium hot side temperature [C]"], abs=1e-9
    )
    heater = f.units["Heater TwinLink"]
    m = with_heater(si.plant.material, heater.parameters["Efficiency"][0] / 100, 55e6)
    burner = (
        Line(LineMode.partner())
        .burner_unit(m)
        .solve(BurnerInputs(c.Stream(gas_dry, gas_y, stage.gas_out_celsius), 250.0))
    )
    assert burner.duty_w / 1e3 == pytest.approx(heater.state["Heat duty [kW]"], rel=2e-3)


# ---------------------------------------------------------------------- validators
def test_validator_hits_on_the_reported_values(sheets, new_inputs) -> None:  # type: ignore[no-untyped-def]
    dr, hr, co = sheets["Dryer"], sheets["Heat recovery"], sheets["Combined"]
    sat = SaturationProps.from_material(new_inputs.to_si().plant.material)
    assert sat.cp_vapour_j_per_kg_k == 1906.0
    saturation, product = set(), set()
    for k in DATA_SETS:
        d = _dryer(dr, k)
        for src in ("excel", "dyssol"):
            m_p = getattr(dr.results["product_kg_per_h"][k], src) / 3600
            t = getattr(dr.results["exhaust_celsius"][k], src)
            y = (
                d.air_humidity_kg_per_kg
                + (d.spray_mass_flow_kg_per_s - m_p) / d.air_mass_flow_kg_per_s
            )
            if not check_exhaust_saturation(t, y).ok:
                saturation.add((k, src))
            if not check_product_mass(
                m_p, d.spray_mass_flow_kg_per_s, d.spray_solid_fraction_kg_per_kg
            ).ok:
                product.add((k, src))
    assert saturation == {(3, "excel"), (3, "dyssol")}
    assert product == {(1, "dyssol"), (2, "dyssol")}
    boiling = {
        (sheet.name, k)
        for sheet in (hr, co)
        for k in DATA_SETS
        if not check_medium_boiling(sheet.results["medium_hot_celsius"][k].excel).ok
    }
    assert boiling == {("Combined", 4), ("Combined", 5), ("Heat recovery", 5)}


def test_physical_mode_removes_the_saturation_violation_and_caps_the_medium(
    sheets, new_inputs
) -> None:  # type: ignore[no-untyped-def]
    for k in DATA_SETS:
        r = _combined(sheets, new_inputs, k, LineMode.physical(htm_max_celsius=95.0))
        p = _combined(sheets, new_inputs, k, LineMode.partner())
        assert r.medium_hot_celsius <= 95.0 + 1e-9
        assert r.recovered_kw >= p.recovered_kw
        assert {f.check for f in r.findings() if not f.ok} == set()
        assert not math.isnan(r.fuel_kg_per_h)


# ---------------------------------------------------------------------- reader (synthetic)
def test_comparison_reader_on_a_synthetic_workbook(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in ("Dryer", "Heater", "Heat recovery", "Combined"):
        ws = book.create_sheet(name)
        ws["A3"], ws["B3"], ws["C3"] = "Unit", "Process air", "Mass flow rate"
        ws["C4"] = "Temperature"
        ws["B5"], ws["C5"] = "Spray rate", "Solid fraction"
        ws["A6"], ws["B6"] = "Heater", "Efficiency"
        ws["A7"], ws["C7"] = "Heating transfer medium", "Flow rate"
        for k in DATA_SETS:
            col = 4 * k  # 0-based value column of data set k (E, I, M, Q, U)
            for row, unit, value in (
                (3, "kg/hr", 100 + k),
                (4, "\N{DEGREE SIGN}C", 200 + k),
                (5, "-", 0.1 * k),
                (6, "%", 70 + k),
                (7, "kg/hr", 10 * k),
            ):
                ws.cell(row, col, unit)
                ws.cell(row, col + 1, value)
        ws["A11"] = "Heat recovered (kW)"
        for k in DATA_SETS:
            ws.cell(11 + k, 1, f"Data set {k}")
            ws.cell(11 + k, 3, 1.0 + k)
            ws.cell(11 + k, 4, 2.0 + k)
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    sheets = load_comparison(path)
    s = sheets["Combined"]
    assert s.value(3, "Process air", "Mass flow rate") == 103
    assert s.value(3, "Process air", "Temperature") == 203
    assert s.value(2, "Spray rate", "Solid fraction") == pytest.approx(0.2)
    assert s.value(5, "Heater", "Efficiency") == 75
    assert s.value(4, "Heating transfer medium", "Flow rate") == 40
    assert s.results["recovered_kw"][2].excel == 3.0
    assert s.results["recovered_kw"][2].dyssol == 4.0
    with pytest.raises(KeyError):
        s.value(1, "Fresh gas", "Temperature")
